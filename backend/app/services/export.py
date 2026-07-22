import io
import re
import unicodedata
from datetime import date
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session
from weasyprint import HTML

from app.models.client import Client
from app.models.intervention import Intervention
from app.schemas.dashboard import NiveauConso
from app.services.consommation import calculer_pourcentage, consommation

TEMPLATES_DIR = Path(__file__).resolve().parent.parent / "templates"
_env = Environment(loader=FileSystemLoader(TEMPLATES_DIR), autoescape=select_autoescape())


def nom_fichier_sur(nom: str) -> str:
    ascii_nom = unicodedata.normalize("NFKD", nom).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Za-z0-9_-]+", "_", ascii_nom).strip("_") or "client"


def construire_rapport(db: Session, client: Client, periode: tuple[date, date]) -> dict:
    periode_debut, periode_fin = periode

    minutes_n1 = consommation(db, client.id, "N1", periode)
    minutes_n2 = consommation(db, client.id, "N2", periode)
    minutes_n3 = consommation(db, client.id, "N3", periode)

    pct_n1 = calculer_pourcentage(minutes_n1, client.forfait_n1_h)
    pct_n2 = calculer_pourcentage(minutes_n2, client.forfait_n2_h)

    heures_supp_minutes = max(0, minutes_n1 - client.forfait_n1_h * 60) + max(
        0, minutes_n2 - client.forfait_n2_h * 60
    )

    interventions = (
        db.query(Intervention)
        .filter(
            Intervention.client_id == client.id,
            Intervention.date_intervention >= periode_debut,
            Intervention.date_intervention <= periode_fin,
        )
        .order_by(Intervention.date_intervention, Intervention.id)
        .all()
    )

    return {
        "client": client,
        "periode_debut": periode_debut,
        "periode_fin": periode_fin,
        "n1": NiveauConso(forfait_h=client.forfait_n1_h, consomme_minutes=minutes_n1, pourcentage=round(pct_n1, 1)),
        "n2": NiveauConso(forfait_h=client.forfait_n2_h, consomme_minutes=minutes_n2, pourcentage=round(pct_n2, 1)),
        "n3": NiveauConso(forfait_h=None, consomme_minutes=minutes_n3, pourcentage=None),
        "heures_supp_minutes": heures_supp_minutes,
        "interventions": interventions,
    }


def generer_pdf(rapport: dict) -> bytes:
    template = _env.get_template("rapport.html")
    html_content = template.render(**rapport)
    return HTML(string=html_content).write_pdf()


def generer_excel(rapport: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapport"

    ws.append([f"Rapport — {rapport['client'].nom}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Période du {rapport['periode_debut']} au {rapport['periode_fin']}"])
    ws.append([])
    ws.append(["Niveau", "Forfait (h)", "Consommé (h)", "%"])
    ws.append(["N1 — Assistance", rapport["n1"].forfait_h, round(rapport["n1"].consomme_minutes / 60, 2), rapport["n1"].pourcentage])
    ws.append(["N2 — Optimisation", rapport["n2"].forfait_h, round(rapport["n2"].consomme_minutes / 60, 2), rapport["n2"].pourcentage])
    ws.append(["N3 — Conseil (hors forfait)", "-", round(rapport["n3"].consomme_minutes / 60, 2), "-"])
    ws.append(["Heures supp.", "-", round(rapport["heures_supp_minutes"] / 60, 2), "-"])
    ws.append([])

    ws.append(["Date", "Niveau", "Durée (min)", "Description"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    for intervention in rapport["interventions"]:
        ws.append(
            [
                intervention.date_intervention.isoformat(),
                intervention.niveau,
                intervention.duree_minutes,
                intervention.description,
            ]
        )

    for col, width in zip("ABCD", (28, 14, 14, 50)):
        ws.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
